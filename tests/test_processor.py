import unittest

from docx import Document
from docx.enum.section import WD_SECTION_START

from processor import (
    QUOTE_LANGUAGES,
    QUOTE_STYLES,
    apply_quote_style_to_segments,
    apply_quote_style_to_text,
    delete_empty_paragraphs,
    normalize_run_text,
    quote_example,
)
from audit_log import AuditLog
from processor import _run_fix_broken_sentences, fix_broken_sentences, remove_page_frames
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class NormalizeRunTextTests(unittest.TestCase):
    def test_keeps_outer_spaces_around_quotes(self):
        text = 'strankama "plavih" i "zelenih"'
        self.assertEqual(normalize_run_text(text), text)

    def test_removes_inner_spaces_inside_quotes(self):
        text = 'strankama " plavih " i " zelenih "'
        expected = 'strankama "plavih" i "zelenih"'
        self.assertEqual(normalize_run_text(text), expected)

    def test_fixes_space_before_comma_and_period(self):
        text = "Ovo je test , a ovo druga recenica ."
        expected = "Ovo je test, a ovo druga recenica."
        self.assertEqual(normalize_run_text(text), expected)

    def test_joins_broken_hyphenated_word(self):
        self.assertEqual(normalize_run_text("pro- gram"), "program")
        self.assertEqual(normalize_run_text("pro - gram"), "program")

    def test_removes_false_spacing_in_numbers(self):
        self.assertEqual(normalize_run_text("1 000 i 12 345"), "1000 i 12345")

    def test_normalizes_ligatures(self):
        self.assertEqual(normalize_run_text("ofﬁce ﬂow"), "office flow")

    def test_normalizes_pdf_quotes(self):
        self.assertEqual(normalize_run_text("``tekst''"), '"tekst"')
        self.assertEqual(normalize_run_text("„tekst”"), '"tekst"')
        self.assertEqual(normalize_run_text("„razbojnik11"), '"razbojnik"')

    def test_does_not_turn_regular_number_11_into_quote(self):
        self.assertEqual(normalize_run_text('clan 11 stupa na snagu'), "clan 11 stupa na snagu")

    def test_fixes_ocr_quote_11_and_missing_space_before_quote(self):
        text = 'Jer nazivi "ubica" i "razbojnik11 bili su od njih cenjeni i odgovarali su nazivu"energičan"'
        expected = 'Jer nazivi "ubica" i "razbojnik" bili su od njih cenjeni i odgovarali su nazivu "energičan"'
        self.assertEqual(normalize_run_text(text), expected)

    def test_normalizes_duplicate_punctuation(self):
        self.assertEqual(normalize_run_text("ovo .. ,,, test"), "ovo., test")

    def test_keeps_an_ellipsis(self):
        self.assertEqual(
            normalize_run_text("trebalo je imati vere... vere koju"),
            "trebalo je imati vere... vere koju",
        )

    def test_tidies_a_spaced_or_overlong_ellipsis(self):
        self.assertEqual(normalize_run_text("vere. . . vere"), "vere... vere")
        self.assertEqual(normalize_run_text("vere..... vere"), "vere... vere")

    def test_keeps_an_en_dash_between_words(self):
        self.assertEqual(
            normalize_run_text("znamo da to mozemo – kao sto znamo"),
            "znamo da to mozemo – kao sto znamo",
        )

    def test_reduces_an_em_dash_to_an_en_dash(self):
        self.assertEqual(normalize_run_text("mozemo — kao sto"), "mozemo – kao sto")
        self.assertEqual(normalize_run_text("— Zdravo"), "– Zdravo")

    def test_spaces_out_a_dash_glued_between_words(self):
        self.assertEqual(normalize_run_text("rec—rec"), "rec – rec")

    def test_slash_spacing_is_off_unless_asked_for(self):
        self.assertEqual(normalize_run_text("i / ili"), "i / ili")

    def test_slash_spacing_closes_up_when_asked_for(self):
        self.assertEqual(normalize_run_text("i / ili", close_slash_spacing=True), "i/ili")

    def test_uniform_quote_style_english(self):
        self.assertEqual(apply_quote_style_to_text('"Crvena zvezda"', "english-double"), '"Crvena zvezda"')

    def test_uniform_quote_style_english_single(self):
        self.assertEqual(apply_quote_style_to_text('"Crvena zvezda"', "english-single"), "'Crvena zvezda'")

    def test_uniform_quote_style_serbian(self):
        self.assertEqual(apply_quote_style_to_text('"Crvena zvezda"', "serbian"), "„Crvena zvezda”")

    def test_uniform_quote_style_german(self):
        self.assertEqual(apply_quote_style_to_text('"Crvena zvezda"', "german"), "„Crvena zvezda“")

    def test_uniform_quote_style_serbian_across_segments(self):
        segments = ['"', "crveni", '"', ", ", '"', "zeleni", '"']
        expected = ["„", "crveni", "”", ", ", "„", "zeleni", "”"]
        self.assertEqual(apply_quote_style_to_segments(segments, "serbian"), expected)


class DeleteEmptyParagraphsTests(unittest.TestCase):
    @staticmethod
    def _clean(doc):
        return delete_empty_paragraphs(doc, lambda message: None)

    def test_keeps_last_paragraph_in_otherwise_empty_table_cell(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "sadrzaj"
        table.cell(0, 1).text = ""

        self._clean(doc)

        self.assertEqual(len(table.cell(0, 1).paragraphs), 1)

    def test_removes_extra_blank_paragraphs_inside_table_cell(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.text = ""
        cell.add_paragraph("")
        cell.add_paragraph("")

        removed = self._clean(doc)

        self.assertEqual(len(cell.paragraphs), 1)
        self.assertEqual(removed, 2)

    def test_keeps_a_blank_paragraph_that_carries_a_section_break(self):
        doc = Document()
        doc.add_paragraph("prvi")
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_paragraph("drugi")

        self._clean(doc)

        self.assertEqual(len(doc.sections), 2)
        self.assertEqual([p.text for p in doc.paragraphs if p.text.strip()], ["prvi", "drugi"])

    def test_removes_all_blank_paragraphs_from_document_body(self):
        doc = Document()
        doc.add_paragraph("prvi")
        doc.add_paragraph("")
        doc.add_paragraph("")
        doc.add_paragraph("drugi")

        removed = self._clean(doc)

        self.assertEqual([p.text for p in doc.paragraphs], ["prvi", "drugi"])
        self.assertEqual(removed, 2)


class MergeAuditTests(unittest.TestCase):
    @staticmethod
    def _audit():
        return AuditLog("src.docx", "dst.docx", {}, "serbian", ".docx", {})

    def test_records_one_change_per_merge_and_nothing_else(self):
        doc = Document()
        doc.add_paragraph("Prva recenica koja se nastavlja")
        doc.add_paragraph("nastavak prve recenice.")
        doc.add_paragraph("Druga recenica koja se nastavlja")
        doc.add_paragraph("nastavak druge recenice.")
        doc.add_paragraph("Peta recenica stoji sama.")
        doc.add_paragraph("Sesta recenica stoji sama.")
        audit = self._audit()

        _run_fix_broken_sentences(doc, lambda message: None, False, None, None, audit)

        self.assertEqual(audit.stats["merged_paragraph_pairs"], 2)
        self.assertEqual(len(audit.changes["paragraph_merge"]), 2)

    def test_recorded_change_shows_the_two_merged_paragraphs(self):
        doc = Document()
        doc.add_paragraph("Prva recenica koja se nastavlja")
        doc.add_paragraph("nastavak prve recenice.")
        audit = self._audit()

        _run_fix_broken_sentences(doc, lambda message: None, False, None, None, audit)

        change = audit.changes["paragraph_merge"][0]
        self.assertIn("Prva recenica koja se nastavlja", change["before"])
        self.assertIn("nastavak prve recenice.", change["before"])
        self.assertEqual(
            change["after"],
            "Prva recenica koja se nastavlja nastavak prve recenice.",
        )


class MergeAcrossSectionBreakTests(unittest.TestCase):
    @staticmethod
    def _merge(doc):
        return fix_broken_sentences(doc, lambda message: None)

    def test_merges_a_sentence_split_across_a_section_break(self):
        doc = Document()
        doc.add_paragraph("Prva recenica koja se nastavlja")
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_paragraph("nastavak prve recenice.")

        merges = self._merge(doc)

        self.assertEqual(merges, 1)
        self.assertIn(
            "Prva recenica koja se nastavlja nastavak prve recenice.",
            [p.text for p in doc.paragraphs],
        )

    def test_the_section_break_survives_the_merge(self):
        doc = Document()
        doc.add_paragraph("Prva recenica koja se nastavlja")
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_paragraph("nastavak prve recenice.")

        self._merge(doc)

        self.assertEqual(len(doc.sections), 2)

    def test_does_not_merge_across_a_plain_blank_paragraph(self):
        doc = Document()
        doc.add_paragraph("Prva recenica koja se nastavlja")
        doc.add_paragraph("")
        doc.add_paragraph("nastavak prve recenice.")

        self.assertEqual(self._merge(doc), 0)


def anchor_to_page(paragraph, x=2754, y=5754):
    """Pin a paragraph to an absolute spot on its page, the way ABBYY converts a PDF."""
    frame = OxmlElement("w:framePr")
    frame.set(qn("w:wrap"), "none")
    frame.set(qn("w:vAnchor"), "page")
    frame.set(qn("w:hAnchor"), "page")
    frame.set(qn("w:x"), str(x))
    frame.set(qn("w:y"), str(y))
    paragraph._element.get_or_add_pPr().insert(0, frame)
    return paragraph


class PageFrameTests(unittest.TestCase):
    @staticmethod
    def _strip(doc):
        return remove_page_frames(doc, lambda message: None)

    @staticmethod
    def _frames(paragraph):
        return paragraph._element.xpath("./w:pPr/w:framePr")

    def test_releases_a_paragraph_pinned_to_the_page(self):
        doc = Document()
        paragraph = anchor_to_page(doc.add_paragraph("tekst"))

        self._strip(doc)

        self.assertEqual(self._frames(paragraph), [])

    def test_releases_paragraphs_inside_table_cells(self):
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "u tabeli"
        paragraph = anchor_to_page(cell.paragraphs[0])

        self._strip(doc)

        self.assertEqual(self._frames(paragraph), [])

    def test_reports_how_many_were_released(self):
        doc = Document()
        anchor_to_page(doc.add_paragraph("prvi"))
        anchor_to_page(doc.add_paragraph("drugi"))
        doc.add_paragraph("treci bez okvira")

        self.assertEqual(self._strip(doc), 2)

    def test_releases_frames_in_vertically_merged_table_cells(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).merge(table.cell(1, 0))
        for element in table._tbl.iter(qn("w:p")):
            frame = OxmlElement("w:framePr")
            frame.set(qn("w:vAnchor"), "page")
            frame.set(qn("w:y"), "5754")
            element.get_or_add_pPr().insert(0, frame)

        self._strip(doc)

        self.assertEqual(list(doc.element.body.iter(qn("w:framePr"))), [])

    def test_leaves_the_text_alone(self):
        doc = Document()
        anchor_to_page(doc.add_paragraph("tekst ostaje isti"))

        self._strip(doc)

        self.assertEqual([p.text for p in doc.paragraphs], ["tekst ostaje isti"])


class QuoteExampleTests(unittest.TestCase):
    def test_wraps_the_sample_in_the_style_of_the_language(self):
        self.assertEqual(quote_example("english-double", "Proxima"), '"Proxima"')
        self.assertEqual(quote_example("english-single", "Proxima"), "'Proxima'")
        self.assertEqual(quote_example("serbian", "Proxima"), "\u201eProxima\u201d")
        self.assertEqual(quote_example("german", "Proxima"), "\u201eProxima\u201c")

    def test_falls_back_to_english_double_for_an_unknown_language(self):
        self.assertEqual(quote_example("klingon", "Proxima"), '"Proxima"')

    def test_every_offered_language_has_a_style(self):
        for language in QUOTE_LANGUAGES:
            self.assertIn(language, QUOTE_STYLES)

    def test_every_offered_language_reads_differently(self):
        examples = [quote_example(language) for language in QUOTE_LANGUAGES]
        self.assertEqual(len(examples), len(set(examples)))


class LegalNumberingTests(unittest.TestCase):
    @staticmethod
    def _merge(doc, protect_legal_numbering):
        return fix_broken_sentences(doc, lambda message: None, protect_legal_numbering)

    @staticmethod
    def _numbered_pair(doc):
        doc.add_paragraph("Article 12 the parties agree that")
        doc.add_paragraph("the contract shall remain in force.")

    def test_merges_numbered_clauses_when_protection_is_off(self):
        doc = Document()
        self._numbered_pair(doc)

        self.assertEqual(self._merge(doc, False), 1)

    def test_leaves_numbered_clauses_alone_when_protection_is_on(self):
        doc = Document()
        self._numbered_pair(doc)

        self.assertEqual(self._merge(doc, True), 0)

    def test_protects_a_label_line_whatever_the_setting(self):
        for protect in (False, True):
            doc = Document()
            doc.add_paragraph("Napomena:")
            doc.add_paragraph("nastavak koji bi se inace spojio.")

            self.assertEqual(self._merge(doc, protect), 0)


if __name__ == "__main__":
    unittest.main()
